"""Load Telegram-export Excel files into normalized message dicts."""

from __future__ import annotations

from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from typing import Iterable, Iterator

import openpyxl


EXPECTED_COLUMNS = [
    "message_id", "date", "text", "group_id",
    "user_id", "username", "first_name", "last_name", "phone",
]


@dataclass
class Message:
    message_id: int | None
    date: datetime | None
    text: str | None
    group_id: int | None
    user_id: int | None
    username: str | None
    first_name: str | None
    last_name: str | None

    def to_dict(self) -> dict:
        d = asdict(self)
        if self.date is not None:
            d["date"] = self.date.isoformat()
        return d


def load_messages(
    xlsx_path: str | Path,
    sheet: str | None = None,
    limit: int | None = None,
    only_with_text: bool = True,
) -> list[Message]:
    """Load messages from a Telegram-export workbook.

    Args:
        xlsx_path: path to the .xlsx file.
        sheet: sheet name; defaults to the first sheet.
        limit: cap the number of messages returned (useful for cheap iteration).
        only_with_text: drop messages that have no text body. Media/stickers
            carry no terminology so they're noise for our purposes.
    """
    path = Path(xlsx_path)
    if not path.exists():
        raise FileNotFoundError(f"Input file not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet_name = sheet or wb.sheetnames[0]
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet {sheet_name!r} not found in {path.name}. "
            f"Available: {wb.sheetnames}"
        )
    ws = wb[sheet_name]

    rows: Iterator[tuple] = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []

    # Build a positional map so we tolerate column reorders / missing columns.
    header_norm = [str(c).strip().lower() if c is not None else "" for c in header]
    idx = {col: header_norm.index(col) for col in EXPECTED_COLUMNS if col in header_norm}
    if "text" not in idx or "message_id" not in idx:
        raise ValueError(
            f"Required columns missing from {sheet_name}. "
            f"Expected at least 'message_id' and 'text'. Got: {header_norm}"
        )

    out: list[Message] = []
    for row in rows:
        def get(col: str):
            i = idx.get(col)
            return row[i] if i is not None and i < len(row) else None

        text = get("text")
        if only_with_text and not (isinstance(text, str) and text.strip()):
            continue

        msg = Message(
            message_id=_to_int(get("message_id")),
            date=_to_dt(get("date")),
            text=text if isinstance(text, str) else None,
            group_id=_to_int(get("group_id")),
            user_id=_to_int(get("user_id")),
            username=_to_str(get("username")),
            first_name=_to_str(get("first_name")),
            last_name=_to_str(get("last_name")),
        )
        out.append(msg)
        if limit is not None and len(out) >= limit:
            break

    wb.close()
    return out


def _to_int(v) -> int | None:
    if v is None:
        return None
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _to_str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _to_dt(v) -> datetime | None:
    if isinstance(v, datetime):
        return v
    return None


def corpus_stats(messages: Iterable[Message]) -> dict:
    """Lightweight summary, mostly for logging."""
    msgs = list(messages)
    if not msgs:
        return {"count": 0}
    dates = [m.date for m in msgs if m.date is not None]
    return {
        "count": len(msgs),
        "date_min": min(dates).isoformat() if dates else None,
        "date_max": max(dates).isoformat() if dates else None,
        "unique_users": len({m.user_id for m in msgs if m.user_id is not None}),
        "unique_usernames": len({m.username for m in msgs if m.username}),
    }
